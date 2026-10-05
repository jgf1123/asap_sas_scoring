"""Print the official ReadMeFirst description (text + tables) for the given set numbers."""
import sys
import docx

for n in sys.argv[1:]:
    doc = docx.Document(f'data/descriptions/Data Set #{n}--ReadMeFirst.docx')
    print(f'===== SET {n} =====')
    for p in doc.paragraphs:
        if p.text.strip():
            print(p.text)
    for t in doc.tables:
        for r in t.rows:
            print(' | '.join(c.text.strip() for c in r.cells))
